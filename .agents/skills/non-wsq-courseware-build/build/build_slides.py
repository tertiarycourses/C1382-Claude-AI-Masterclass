#!/usr/bin/env python3
"""Build the Claude AI Masterclass deck in the Claude Style house design.

Content comes entirely from course_data.py + data_domain1..3.py, so the deck,
Lesson Plan, Learner Guide and labs are generated from one source and cannot
drift apart. Layout comes from claude_style.py, the component library shared
with the Claude Code Masterclass.
"""
import os, sys, glob, re, importlib
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from claude_style import Deck, CLAY, TEAL, INK


def load_domains():
    acts = []
    files = sorted(glob.glob(os.path.join(HERE, "data_domain[0-9]*.py")),
                   key=lambda p: int("".join(c for c in os.path.basename(p) if c.isdigit()) or 0))
    for f in files:
        if not re.fullmatch(r"data_domain\d+\.py", os.path.basename(f)):
            continue
        mod = importlib.import_module(os.path.basename(f)[:-3])
        n = "".join(c for c in os.path.basename(f) if c.isdigit())
        acts += getattr(mod, f"DOMAIN{n}", [])
    return acts


ACTIVITIES = load_domains()
# Spelled-out lab count, so prose never disagrees with the data.
_NUMWORDS = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six",
             7: "Seven", 8: "Eight", 9: "Nine", 10: "Ten"}
_LABWORD = _NUMWORDS.get(len(ACTIVITIES), str(len(ACTIVITIES)))
TOPIC_ACTS = {t["num"]: [a for a in ACTIVITIES if a["topic"] == t["num"]] for t in C.TOPICS}


def find_repo(start):
    env = os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env):
        return env
    d = start
    for _ in range(8):
        d = os.path.dirname(d)
        if os.path.isdir(os.path.join(d, "courseware")) and os.path.isdir(os.path.join(d, "labs")):
            return d
    return os.path.dirname(os.path.dirname(start))


REPO = find_repo(HERE)
# Lab output screenshots, referenced by a lab's `verify_shot` filename.
SHOTS = os.path.join(REPO, "courseware", "assets", "screenshots")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
D = Deck(prs, C)

# ============================================================== COVER + ADMIN
D.cover()

D.section("COURSE ADMINISTRATION", "Welcome", ["Introductions", "Ground rules", "Course outline", "How today runs"])

D.tiles("Let's Know Each Other", C.ICE_BREAKER, kicker="ICE-BREAKER",
        intro="Take a minute to introduce yourself to the class.", cols=2)

# Trainer profiles: the blank template the delivering trainer completes, then
# Alfred's. Both follow the learner self-introduction.
D.trainer(C.TRAINER_BLANK)
D.trainer(C.TRAINER_PROFILE)

D.tiles("Ground Rules", C.GROUND_RULES, kicker="HOUSEKEEPING", cols=2, numbered=False)

D.outline("Course Outline", [(t["title"], t["agenda"]) for t in C.TOPICS])

# Learning outcomes
D.tiles("Learning Outcomes",
        [(C.LO_TITLES[i], lo.split(": ", 1)[1] if ": " in lo else lo)
         for i, lo in enumerate(C.LEARNING_OUTCOMES)],
        kicker="WHAT YOU'LL ACHIEVE", cols=2)

# Lesson plan — built from the same SCHEDULE the LP uses
def _fmt_time(hhmm):
    h, m = (int(x) for x in hhmm.split(":"))
    ap = "am" if h < 12 else "pm"
    hh = h if 1 <= h <= 12 else (h - 12 if h > 12 else 12)
    return f"{hh}:{m:02d}{ap}"


_lab = {a["num"]: a["title"] for a in ACTIVITIES}
_sched = C.SCHEDULE(lambda ns: " · ".join(_lab[n] for n in ns if n in _lab))
_rows = _sched[1][1]


def _mins(hhmm):
    """'9:00' -> 540. Compare times numerically: as strings '9:00' sorts AFTER
    '12:30', which silently put the morning rows in the afternoon column."""
    h, m = (int(x) for x in hhmm.split(":"))
    return h * 60 + m


# Split at lunch, so each column is a real half-day rather than a fixed clock time.
_lunch_at = next((_mins(r[0]) for r in _rows if r[3] == "lunch"), 12 * 60 + 30)


def _row_text(r):
    """One informative line per entry. The full schedule text is written for the
    Lesson Plan's wide table, so shorten it here to what the narrow column can
    show — the topic name, or the lab numbers — rather than a bare 'Topic 1'."""
    kind, text = r[3], r[4]
    if kind == "topic":
        # "Topic 2 — Tools and Skills: connectors, …" -> "Topic 2 — Tools and Skills"
        label = text.split(":")[0].strip()
    elif kind == "lab":
        # The schedule row lists lab TITLES; map them back to numbers so the
        # narrow column can show "Hands-on: Labs 1, 2" instead of a long list.
        nums = [str(a["num"]) for a in ACTIVITIES if a["title"] in text]
        label = (("Hands-on: Lab" + ("s " if len(nums) > 1 else " ") + ", ".join(nums))
                 if nums else "Hands-on lab")
    else:
        label = text.split(",")[0].strip()
    return f"{r[0]}  {label[:56]}"


_morning = [_row_text(r) for r in _rows if _mins(r[0]) < _lunch_at]
_afternoon = [_row_text(r) for r in _rows if _mins(r[0]) >= _lunch_at]
D.split("Lesson Plan — One Day", ("Morning", _morning), ("Afternoon", _afternoon),
        kicker="SCHEDULE",
        intro=f"{_fmt_time(_rows[0][0])} – {_fmt_time(_rows[-1][1])}, with a one-hour lunch and two tea breaks.")

D.flow("How You'll Learn",
       [("Watch", "The trainer demonstrates the technique live."),
        ("Build", "You build it yourself in the hands-on lab."),
        ("Check", "Verify your output against the lab's check step."),
        ("Discuss", "We troubleshoot together as a class."),
        ("Reuse", "Take the artefact and the prompt back to your own work.")],
       kicker="LEARNING APPROACH",
       foot=f"{_LABWORD} labs, one continuous scenario — each lab builds on the one before it.")

# Where the copy-ready prompts live. A light tiles slide rather than a second
# dark statement, so it does not sit back to back with the scenario slide.
D.tiles("Your Course Materials",
        [("Copy-ready prompts", f"{C.PROMPTS_URL} — every lab's prompt with a copy button. "
                                "Keep it open in a browser tab all day."),
         ("Lab files and data", f"{C.REPO_URL} — each lab folder holds its prompt PDF and, "
                                "where needed, a data folder of mock files."),
         ("Your Learner Guide", "The full step-by-step for every lab, the reason for each step, "
                                "and the check to run at the end.")],
        kicker="BEFORE WE START", cols=1,
        intro="Open the prompts page now — you will paste from it in every lab.")

# The scenario that runs through every lab, in two slides: the situation on
# the dark statement slide, then the learner's role and mission. The synthetic-
# data disclaimer is the last sentence of COMPANY_BLURB, so it is taken from
# there rather than restated here.
_SC = C.COURSE_SCENARIO
D.statement(_SC["title"], _SC["situation"], kicker="THE SCENARIO")
_disclaimer = re.split(r"(?<=\.)\s+", C.COMPANY_BLURB.strip())[-1]
D.tiles(f"One Brand, {_LABWORD} Labs, One Project",
        [("Your role", _SC["role"]),
         ("Your mission", _SC["mission"]),
         ("Synthetic data only", _disclaimer)],
        kicker="THE SCENARIO", cols=1, numbered=False)

# Theory that opens each topic, before its key concepts and first lab. Theory
# tied to ONE lab is not here — it sits in BEFORE_LAB, next to the lab.
TOPIC_INTRO = {
    1: (
        lambda: D.cards("The Claude Product Family", C.CLAUDE_PRODUCTS,
                        kicker="TOPIC 1 · OVERVIEW",
                        intro="Anthropic offers several Claude products — here's how they differ so you know which is which.",
                        cols=3, tip=C.CLAUDE_PRODUCTS_TIP),
        lambda: D.compare("Claude (Chat + Cowork) and Claude Code", C.CLAUDE_THREE,
                          kicker="TOPIC 1 · OVERVIEW", foot=C.CLAUDE_THREE_FOOT),
        lambda: D.timeline("The Evolution of AI Engineering", C.AI_EVOLUTION,
                           kicker="TOPIC 1 · OVERVIEW", foot=C.AI_EVOLUTION_FOOT),
        lambda: D.tiles("What Cowork Mode Adds", C.COWORK_CAPABILITIES,
                        kicker="TOPIC 1 · OVERVIEW", cols=2,
                        intro="Five things a chat window cannot do.",
                        foot=C.COWORK_CAPABILITIES_FOOT),
        lambda: D.pricing("Claude Pricing", C.PRICING_INTRO, C.PRICING_PLANS,
                          C.PRICING_BAND, C.PRICING_TIP, kicker="TOPIC 1 · SETUP"),
        lambda: D.tiles("Set Up Claude", C.COWORK_SETUP, kicker="TOPIC 1 · SETUP",
                        intro="Projects, Cowork mode and Live Artefacts run in the Claude desktop app on a paid plan.",
                        cols=2),
        # ---- the redesigned Claude Projects: what they are, then how to set one up.
        # Lab 1 creates the project every later lab runs inside, so all of this
        # comes before it.
        lambda: D.split("The New Claude Projects",
                        C.PROJECTS_BEFORE_AFTER["left"], C.PROJECTS_BEFORE_AFTER["right"],
                        kicker="TOPIC 1 · PROJECTS",
                        intro="You stop managing chats and start briefing a team."),
        lambda: D.container("A Project Is a Team", "Claude Project", C.PROJECT_TEAM_SUBTITLE,
                            C.PROJECT_TEAM_PARTS, kicker="TOPIC 1 · LEAD AGENT AND THREADS",
                            foot=C.PROJECT_TEAM_FOOT),
        lambda: D.tiles("When to Use a Project", C.PROJECT_WHEN, kicker="TOPIC 1 · PROJECTS",
                        cols=2, foot=C.PROJECT_WHEN_FOOT),
        lambda: D.tiles("Anatomy of a Project", C.PROJECT_ANATOMY, kicker="TOPIC 1 · PROJECTS",
                        intro="Six parts — the first four you set up in Lab 1, the last two run the work.",
                        cols=2),
        lambda: D.numbered_rows("Creating a Project, Field by Field", C.PROJECT_DIALOG,
                                kicker="TOPIC 1 · PROJECTS",
                                intro="New project opens the Create a project dialog.",
                                foot=C.PROJECT_DIALOG_FOOT),
        lambda: D.numbered_rows("Project Settings", C.PROJECT_SETTINGS,
                                kicker="TOPIC 1 · PROJECTS", foot=C.PROJECT_SETTINGS_FOOT),
        lambda: D.table("Models and Effort", C.PROJECT_MODELS_HEAD, C.PROJECT_MODELS_ROWS,
                        kicker="TOPIC 1 · PROJECT SETTINGS",
                        intro="The lead agent and the threads each have their own model and effort.",
                        foot=C.PROJECT_MODELS_FOOT, highlight_last=False),
        lambda: D.table("Instructions, Context or Memory?", C.MEMORY_CONTEXT_HEAD,
                        C.MEMORY_CONTEXT_ROWS, kicker="TOPIC 1 · PROJECTS",
                        intro="When the output is wrong, this tells you which one to edit.",
                        highlight_last=False),
        lambda: D.split("Why a Project?", C.PROJECT_VS_CHAT["left"], C.PROJECT_VS_CHAT["right"],
                        kicker="TOPIC 1 · PROJECTS",
                        intro="A dedicated place for ongoing work, where context builds over time."),
    ),
    2: (
        lambda: D.tiles("What Connectors Do", C.CONNECTOR_CAPABILITIES, kicker="TOPIC 2 · CONNECTORS",
                        intro="Connectors let Claude reach data and services that live outside the chat.",
                        cols=2),
        lambda: D.flow("Turning On a Connector", C.CONNECTOR_STEPS, kicker="TOPIC 2 · CONNECTORS",
                       foot="Manage them any time under Customize → Connectors."),
        lambda: D.quote("Context Window", C.CONTEXT_WINDOW_QUOTE, kicker="TOPIC 2 · CONTEXT",
                        foot="Files, instructions, tools and skills all have to fit inside it."),
        lambda: D.meter("What Fills the Context Window", C.CONTEXT_WINDOW_SEGMENTS,
                        kicker="TOPIC 2 · CONTEXT",
                        intro="At the start of a thread.", foot=C.CONTEXT_WINDOW_FOOT),
        lambda: D.meter("The Same Window, Later in the Thread", C.CONTEXT_WINDOW_SEGMENTS_FULL,
                        kicker="TOPIC 2 · CONTEXT",
                        intro="After reading files, calling tools and running skills.",
                        foot=C.CONTEXT_WINDOW_FULL_FOOT),
        lambda: D.dark_list("When the Context Window Fills Up", C.CONTEXT_ROT,
                            kicker="TOPIC 2 · CONTEXT", foot=C.CONTEXT_ROT_FOOT),
        lambda: D.numbered_rows("Five Levers of Context", C.CONTEXT_LEVERS,
                                kicker="TOPIC 2 · CONTEXT",
                                intro="What you actually control.", foot=C.CONTEXT_LEVERS_FOOT),
        lambda: D.table("Check It Yourself: /context", C.CONTEXT_COMMAND_HEAD,
                        C.CONTEXT_COMMAND_ROWS, kicker="TOPIC 2 · CONTEXT",
                        intro="Estimated usage by category, for a 1M-token window.",
                        foot=C.CONTEXT_COMMAND_NOTE),
    ),
    3: (
        lambda: D.flow("One Thread per Channel", C.CONTENT_THREADS_FLOW,
                       kicker="TOPIC 3 · CONTENT IN THREADS",
                       foot="The same pattern as the Lab 2 research — one message, independent subjects — applied to production."),
        lambda: D.tiles("What Each Platform Needs", C.SOCIAL_SPECS, kicker="TOPIC 3 · SOCIAL SPECS",
                        intro="Each thread gets its platform's spec as its brief.",
                        cols=2, foot=C.SOCIAL_SPECS_FOOT),
        lambda: D.tiles("Marketing Responsibly with AI", C.RESPONSIBLE_MARKETING,
                        kicker="TOPIC 3 · GOOD PRACTICE",
                        intro="Fast production is only useful if every post is safe to publish.",
                        cols=2),
    ),
}

# Theory slides tied to ONE lab, emitted immediately before it rather than in a
# block at the top of the topic. Keyed by lab number.
BEFORE_LAB = {
    # Lab 2 is the first job handed to the lead agent to split.
    2: (
        lambda: D.tiles("Market Research in Four Threads", C.RESEARCH_THREADS,
                        kicker="TOPIC 1 · RESEARCH IN THREADS",
                        intro="Four subjects, none depending on another — so they can run at once.",
                        cols=2, foot=C.RESEARCH_FOOT),
    ),
    # Live Artefacts are what Lab 3 builds; Labs 1-2 need none of this, so the
    # theory sits between Lab 2 and Lab 3.
    3: (
        lambda: D.split("Why a Live Artefact?",
                        C.ARTEFACT_VS_CHAT["left"], C.ARTEFACT_VS_CHAT["right"],
                        kicker="TOPIC 1 · LIVE ARTEFACTS",
                        intro="The difference between an answer you read and a tool you keep."),
        lambda: D.flow("Building a Live Artefact", C.LIVE_ARTEFACT_FLOW,
                       kicker="TOPIC 1 · LIVE ARTEFACTS",
                       foot="You describe the outcome in plain English — Claude writes and runs the code."),
        lambda: D.split("What an Artefact Can and Cannot Do",
                        C.ARTEFACT_LIMITS["left"], C.ARTEFACT_LIMITS["right"],
                        kicker="TOPIC 1 · LIVE ARTEFACTS",
                        intro="Artefacts run in a sandbox, so anything needing the internet has to come in another way."),
    ),
    5: (
        lambda: D.tiles("Slides, Docs and Built-in Skills", C.SLIDES_DOCS,
                        kicker="TOPIC 2 · SLIDES AND DOCS",
                        intro="Claude's own formats for the things a launch meeting needs.",
                        cols=2, foot=C.SLIDES_DOCS_FOOT),
    ),
    # Lab 6 writes the first skill, so the skill theory sits right before it.
    6: (
        lambda: D.layers("Skills Are Documented on Three Layers", C.SKILL_LAYERS,
                         kicker="TOPIC 2 · SKILLS",
                         intro="Each layer reaches Claude at a different moment — which is why the description matters most."),
        lambda: D.tiles("Anatomy of an Agent Skill", C.SKILL_ANATOMY, kicker="TOPIC 2 · SKILLS",
                        intro="A skill is a folder Claude loads on demand when the task matches its description.",
                        cols=2),
        lambda: D.flow("Creating Your Own Skill", C.SKILL_PROCESS, kicker="TOPIC 2 · SKILLS",
                       intro="Do not write a skill from a blank page — do the work first, then capture it.",
                       foot="/skill-creator writes the SKILL.md for you, from the run you just verified."),
        lambda: D.split("A Skill vs Retyping the Request", C.SKILL_VS_PROMPT["left"],
                        C.SKILL_VS_PROMPT["right"], kicker="TOPIC 2 · SKILLS",
                        intro="Write the routine once, and run it by name from then on."),
        lambda: D.tiles("A Skill That Drives Your Tools", C.SKILL_AUTOMATION,
                        kicker="TOPIC 2 · SKILLS",
                        intro="The most useful skills combine the connectors into one repeatable command.",
                        cols=2),
    ),
    # Lab 7 installs a plugin, so the skill / plugin / mod distinction lands
    # just after the learner has written a skill of their own.
    7: (
        lambda: D.container("Skills vs. Plugins", "Plugin", C.PLUGIN_SUBTITLE, C.PLUGIN_PARTS,
                            kicker="TOPIC 2 · PLUGINS", foot=C.PLUGIN_FOOT),
        lambda: D.tiles("Mods: Changing Claude Itself", C.MODS, kicker="TOPIC 2 · MODS", cols=2),
        lambda: D.table("Skill, Plugin or Mod?", C.SKILLS_PLUGINS_MODS_HEAD,
                        C.SKILLS_PLUGINS_MODS_ROWS, kicker="TOPIC 2 · PLUGINS AND MODS",
                        intro="Three ways to extend Claude — and which one you reach for.",
                        highlight_last=False),
    ),
    9: (
        lambda: D.flow("The Video Pipeline", C.VIDEO_PIPELINE, kicker="TOPIC 3 · LAUNCH VIDEO",
                       foot="The video is code: a change is a re-render, not a re-edit."),
        lambda: D.tiles("Launch Video Specs", C.VIDEO_SPECS, kicker="TOPIC 3 · LAUNCH VIDEO",
                        intro="What the finished 15-second video must meet.", cols=2),
    ),
    10: (
        lambda: D.tiles("Anatomy of a Scheduled Task", C.SCHEDULED_ANATOMY,
                        kicker="TOPIC 3 · SCHEDULED TASKS",
                        intro="A prompt the project runs on a timetable, while you are away.",
                        cols=2),
        lambda: D.numbered_rows("Scheduled Tasks for a Launch", C.SCHEDULED_EXAMPLES,
                                kicker="TOPIC 3 · SCHEDULED TASKS",
                                intro="Each drafts for your approval — none publishes on its own.",
                                foot=C.SCHEDULED_FOOT),
        lambda: D.tiles("Cloud Environments", C.CLOUD_ENVIRONMENT,
                        kicker="TOPIC 3 · SCHEDULED TASKS",
                        intro="Project settings → Environment: let threads and tasks run with your laptop closed.",
                        cols=1),
    ),
}

# Real screenshots of the Claude desktop app, keyed by the lab they lead into.
# A screenshot follows that lab's theory and sits directly before its cover —
# "this is what it looks like in the app, now do it". Only files that exist are
# emitted, so the deck is complete before any capture is made.
SHOTS_BEFORE_LAB = {}
for _key, _file, _title, _cap, _where in C.SCREENSHOTS:
    _path = os.path.join(SHOTS, _file)
    if _where.startswith("lab") and os.path.exists(_path):
        SHOTS_BEFORE_LAB.setdefault(int(_where[3:]), []).append((_path, _title, _cap))



def _pkicker(num, step):
    """Prompt-slide kicker. Naming the step the prompt belongs to lets a
    learner match the slide to the numbered list on the steps slide."""
    return f"LAB {num:02d} · THE PROMPT" + (f" · STEP {step}" if step else "")


# ============================================================== TOPIC LOOP
for t in C.TOPICS:
    acts = TOPIC_ACTS[t["num"]]
    D.section(f"TOPIC {t['code']}", t["title"], t["agenda"])

    for emit in TOPIC_INTRO.get(t["num"], ()):
        emit()

    # ---- key concepts for the topic
    D.tiles(f"Key Concepts — {t['title']}", t["concepts"],
            kicker=f"TOPIC {t['code']} · KEY CONCEPTS", cols=2)

    # ---- labs
    for a in acts:
        # Theory that belongs to one particular lab is emitted immediately
        # before it, so a concept is never taught pages away from the exercise
        # that uses it.
        for emit in BEFORE_LAB.get(a["num"], ()):
            emit()
        for _path, _title, _cap in SHOTS_BEFORE_LAB.get(a["num"], ()):
            D.screenshot(_title, _path, _cap, kicker=f"TOPIC {t['num']} · IN THE APP")
        D.lab_cover(a["num"], a["title"], a["desc"], a["build"], a["services"],
                    objective=a.get("objective"))
        # A multi-message lab gets one prompt slide per block, labelled, so the
        # trainer can stop the class between them for the verify step. Pasting
        # every block at once is the mistake the split exists to prevent, and a
        # single slide holding them all would invite exactly that.
        _plabel = a.get("prompt_label", "PROMPT — copy into Claude")
        if a.get("prompt_blocks"):
            for b in a["prompt_blocks"]:
                # Only the block name goes in the title; the lab's prompt_title
                # is already on the lab cover, and repeating it here makes a
                # heading too long for one line.
                D.prompt(b["label"], b["text"],
                         kicker=_pkicker(a["num"], b.get("step")),
                         intro=b.get("slide_intro"),
                         note=b.get("slide_note"),
                         label=b.get("slide_label", _plabel))
        else:
            D.prompt(a["prompt_title"], a["prompt"],
                     kicker=_pkicker(a["num"], a.get("prompt_step")),
                     note=a.get("prompt_note"), label=_plabel)
        # One summary slide, not one slide per step. The full step-by-step with
        # the "why this step" notes lives in the Learner Guide, which is where
        # someone working through it actually looks. A long lab may supply
        # `slide_steps`: the same sequence condensed to fit one readable slide.
        D.steps_summary(a["num"], a["title"],
                        a.get("slide_steps") or [s[0] for s in a["steps"]])
        # A lab may ship a screenshot of its finished output; it is shown beside
        # the checklist so learners can compare against the real thing.
        _shot = a.get("verify_shot")
        D.verify(a["num"], a["title"], a["test"],
                 shot=(os.path.join(SHOTS, _shot) if _shot else None),
                 caption=a.get("verify_caption"))

    # ---- topic recap
    D.tiles(f"Recap — {t['title']}",
            [(f"Lab {a['num']:02d}", "You can now " + a["objective"] + ".") for a in acts],
            kicker=f"TOPIC {t['code']} · RECAP", cols=1)

# ============================================================== CLOSE
D.section("WRAP-UP", "Summary & Next Steps", ["What you built", "Staying in control", "Next steps", "Q&A"])

D.tiles("What You Achieved",
        [(C.LO_TITLES[i], lo.split(": ", 1)[1] if ": " in lo else lo)
         for i, lo in enumerate(C.LEARNING_OUTCOMES)],
        kicker="LEARNING OUTCOMES", cols=2)

D.tiles("Staying in Control", C.CONTROL_RULES, kicker="GOOD PRACTICE",
        intro="Claude drafts and builds; you check and approve.", cols=2, numbered=False)

D.tiles("Next Steps", C.NEXT_STEPS, kicker="AFTER THE COURSE", cols=2)

D.statement(C.THANK_YOU["headline"], C.THANK_YOU["body"], kicker=C.THANK_YOU["kicker"])

# Fade transition on every slide, matching the Claude Code masterclass deck.
from lxml import etree as _et
from pptx.oxml.ns import qn as _qn
_TRANS = ('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'spd="med"><p:fade/></p:transition>')
for _s in prs.slides:
    if _s.element.find(_qn('p:transition')) is None:
        _s.element.append(_et.fromstring(_TRANS))

OUT = os.path.join(REPO, "courseware", f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
_cp = prs.core_properties
_cp.author = _cp.last_modified_by = C.ORG
_cp.title, _cp.subject, _cp.comments = C.TITLE, f"{C.COURSE_CODE} · {C.VERSION}", ""
prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
