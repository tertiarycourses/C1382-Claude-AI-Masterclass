#!/usr/bin/env python3
"""Generate the non-WSQ Learner Guide as BOTH a Markdown
mirror (LG-<SHORT_TITLE>.md at repo root) and a DOCX (<COURSEWARE_SUBDIR>/LG-*.docx)
from one source, so they never diverge. COURSE_EDITION=classic|threads picks the
edition; a course_data value of None leaves its section out.

House format: cover page, Document Version Control Record, auto TOC, Arial 11pt
body. From v5.0 the lab folders hold only a prompt PDF and a data folder, so this
guide is the ONLY place carrying the full instructions for every lab: goal, what
you build, tools, data files, every prompt in a shaded box, tips, numbered steps
with their reasons, and the check step. Before the topics sits a "Working in a
Claude Project" chapter (goal, instructions, context, memory, scheduled tasks).
Each topic opens with its key concepts, and its theory sections sit immediately
before the lab that uses them — the same order as the day's timetable.

All content is driven by course_data + the domain data files, keeping the LG
100% aligned with the slide deck, Lesson Plan and labs. Screenshots listed in
course_data.SCREENSHOTS are inserted as captioned figures only when the PNG
exists under courseware/assets/screenshots/, so the guide reads fine without them.
"""
import os, sys, re, textwrap
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
# Dynamic domain import — supports ANY number of data_domainN.py files.
import importlib, glob as _g0
def _load_domains():
    acts=[]
    for f in sorted(_g0.glob(os.path.join(HERE,"data_domain[0-9]*.py")),
                    key=lambda q:int("".join(c for c in os.path.basename(q) if c.isdigit()) or 0)):
        n="".join(c for c in os.path.basename(f) if c.isdigit())
        acts+=getattr(importlib.import_module(os.path.basename(f)[:-3]),f"DOMAIN{n}",[])
    return acts
ACT=_load_domains()
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")
SHOT_DIR=os.path.join(REPO,"courseware","assets","screenshots")
SHOT_REL="courseware/assets/screenshots"          # as linked from the repo-root Markdown

def lab_folder(a):
    """<LABS_SUBDIR>/<this>/ — e.g. lab-01-set-up-the-launch-project (same rule as build_labs.py)."""
    slugs=getattr(C,"LAB_SLUGS",{})
    return f"lab-{a['num']:02d}-{slugs[a['num']]}" if a["num"] in slugs else f"lab-{a['num']:02d}"

# ---------------- block DSL (single content stream → MD + DOCX) ----------------
# Text may carry Markdown **bold**; both renderers honour it (the DOCX as real
# bold runs, never as literal asterisks).
B=[]
def h1(t,newpage=False): B.append(("h1",t,newpage))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def h4(t): B.append(("h4",t))                     # small bold label (prompt block, tip title)
def p(t):  B.append(("p",t))
def lead(label,t): B.append(("lead",label,t))     # "Label: text" with a bold label
def bullets(xs): B.append(("bullets",xs))
def defs(pairs,ordered=False): B.append(("defs",pairs,ordered))   # bold term — text
def steps(xs):
    """xs: [(instruction, why-or-None[, attached])]. `attached` is a list of
    prompt dicts (label, lead, text, after, note) shown directly under the step
    where the prompt is sent."""
    B.append(("steps",[(x[0],x[1],x[2] if len(x)>2 else []) for x in xs]))
def note(t): B.append(("note",t))
def check(t): B.append(("check",t))
def table(head,rows,widths,bold_first=True): B.append(("table",head,rows,widths,bold_first))
def rule(): B.append(("rule",))
def pagebreak(): B.append(("pagebreak",))
# A ready-to-paste prompt the learner copies into Claude verbatim. Rendered as a
# fenced text block in Markdown and a shaded monospace panel in the DOCX, so it
# is visually distinct from the surrounding prose.
def prompt(text,header="PROMPT — copy into Claude"): B.append(("prompt",_wrap_prompt(text),header))

def _wrap_prompt(text,width=72):
    """Hard-wrap long lines so a prompt box never relies on the renderer to wrap
    Courier text. Lab prompts are already wrapped; the how-to examples are not."""
    out=[]
    for line in text.strip().split("\n"):
        if len(line)<=width: out.append(line); continue
        ind="  " if line.lstrip().startswith("- ") else ""
        out+=textwrap.wrap(line,width=width,subsequent_indent=ind,break_on_hyphens=False)
    return "\n".join(out)

# ---------------- screenshots ----------------
SHOTS={s[0]:s for s in getattr(C,"SCREENSHOTS",[])}
_FIG={}                                            # key -> figure number once placed
def figure(key):
    """Insert screenshot `key` as a numbered, captioned figure — only if its PNG
    exists. A screenshot already shown earlier becomes a one-line pointer."""
    s=SHOTS.get(key)
    if not s: return
    _k,fname,title,caption,_where=s
    path=os.path.join(SHOT_DIR,fname)
    if not os.path.exists(path): return
    if key in _FIG:
        p(f"See Figure {_FIG[key]} — {title}."); return
    _FIG[key]=len(_FIG)+1
    B.append(("figure",path,fname,_FIG[key],title,caption))
def figures_for(where):
    for s in getattr(C,"SCREENSHOTS",[]):
        if s[4]==where: figure(s[0])

def _oneline(s): return " ".join(s.split())

def compare(d):
    """A before/after or can/cannot pair from course_data as a two-column table.
    A trailing 'So: …' item is a conclusion, not a row, so it follows the table."""
    (lt,litems),(rt,ritems)=d["left"],d["right"]
    tail=[x for x in ritems+litems if x.startswith("So:")]
    li=[x for x in litems if not x.startswith("So:")]; ri=[x for x in ritems if not x.startswith("So:")]
    n=max(len(li),len(ri))
    rows=[(li[i] if i<len(li) else "", ri[i] if i<len(ri) else "") for i in range(n)]
    table((lt,rt),rows,(3.0,3.0),bold_first=False)
    for t_ in tail: p(t_)

# ---------------- front matter ----------------
h1("Introduction")
p(getattr(C,"LG_INTRO",
  f"This Learner Guide accompanies the course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}."))
_SC=getattr(C,"COURSE_SCENARIO",None)
if _SC:
    h3("The scenario" + (f" — {_SC['title']}" if _SC.get("title") else ""))
    for _lbl,_k in (("The situation","situation"),("Your role","role"),("Your mission","mission")):
        if _SC.get(_k): lead(_lbl,_SC[_k])
p(getattr(C,"LG_INTRO2",""))
p(f"This guide carries the full instructions for all {len(ACT)} labs. Each lab folder in the course "
  "repository holds only the lab's prompt PDF and, where the lab needs it, a data folder; the goal, "
  "every prompt, the step-by-step instructions and the check step for each lab are here.")

h1("Course Learning Outcomes")
p("On completion of this course, you will be able to:")
bullets(C.LEARNING_OUTCOMES)

_SETUP=getattr(C,"LG_SETUP",None)
if _SETUP:
    h1("Before You Start — Preparation")
    if _SETUP.get("needs"):
        h3("What you need"); bullets(_SETUP["needs"])
    if _SETUP.get("verify_text") or _SETUP.get("verify_code"):
        h3("Verify your setup")
        if _SETUP.get("verify_text"): p(_SETUP["verify_text"])
    _links=[(lbl,url) for lbl,url in (("Course repository (lab files)",getattr(C,"REPO_URL","")),
                                      ("Copy-ready prompts page",getattr(C,"PROMPTS_URL",""))) if url]
    if _links:
        h3("Course links"); table(("Resource","Address"),_links,(2.1,3.9))
    if _SETUP.get("conventions"):
        h3("Conventions used in every lab"); bullets(_SETUP["conventions"])
    if getattr(C,"CONTROL_RULES",None):
        h3("Rules for every lab"); defs(C.CONTROL_RULES)
        if getattr(C,"LAB_NOTE",""): note(C.LAB_NOTE)

# ---------------- Working in a Claude Project ----------------
h1("Working in a Claude Project",newpage=True)
p("Every lab in this course runs inside one Claude Project — "
  f"{getattr(C,'PROJECT_NAME','the course project')} — which you create in Lab 1. This chapter "
  "explains how a project works and how to set up each of its parts. Come back to it whenever a "
  "lab asks you to change the goal, instructions, context, memory or a scheduled task.")

h2("The lead agent and its threads" if C.THREADS else "How a project is organised")
p(C.PROJECT_TEAM_SUBTITLE + ". " + C.PROJECT_TEAM_FOOT)
defs([(n,_oneline(d)) for n,d,_c in C.PROJECT_TEAM_PARTS])
if getattr(C,"PROJECTS_BEFORE_AFTER",None):
    p("How this changes the way you work:")
    compare(C.PROJECTS_BEFORE_AFTER)

h3("When to use a project")
table(("Use a project for","Why"),C.PROJECT_WHEN,(1.8,4.2))
p(C.PROJECT_WHEN_FOOT)

h2("The parts of a project")
table(("Part","What it is"),C.PROJECT_ANATOMY,(1.5,4.5))
if getattr(C,"PROJECT_DIALOG",None):
    h3("Creating a project")
    defs(C.PROJECT_DIALOG,ordered=True)
    p(C.PROJECT_DIALOG_FOOT)
h3("Instructions, context and memory compared")
p("These three are easy to confuse, and the difference matters: when an output is wrong, it tells you which one to fix.")
table(C.MEMORY_CONTEXT_HEAD,C.MEMORY_CONTEXT_ROWS,(1.5,1.5,1.5,1.5))

for _h in getattr(C,"LG_PROJECT_HOWTO",[]):
    h2(_h["title"])
    if _h.get("shot"): figure(_h["shot"])
    for _p in _h.get("paras",[]): p(_p)
    if _h.get("example"):
        h4("Example")
        prompt(_h["example"],"EXAMPLE — Lumo Kopi")
    if _h.get("steps"):
        h4("How to do it")
        steps([(s_,None) for s_ in _h["steps"]])

# Settings, models and cloud environments exist only in the threads edition;
# in the classic edition these keys are None and the section is left out.
_SETTINGS=getattr(C,"PROJECT_SETTINGS",None); _MODELS=getattr(C,"PROJECT_MODELS_ROWS",None)
_CLOUD=getattr(C,"CLOUD_ENVIRONMENT",None)
if _SETTINGS or _MODELS or _CLOUD:
    h2("Project settings, models and environments")
    if _SETTINGS:
        table(("Settings page","What you find there"),_SETTINGS,(1.5,4.5))
        if C.PROJECT_SETTINGS_FOOT: p(C.PROJECT_SETTINGS_FOOT)
    if _MODELS and C.PROJECT_MODELS_HEAD:
        h3("Models and effort")
        table(C.PROJECT_MODELS_HEAD,_MODELS,(1.1,1.2,1.4,2.3))
        if C.PROJECT_MODELS_FOOT: p(C.PROJECT_MODELS_FOOT)
    if _CLOUD:
        h3("Cloud environments")
        p(C.SCHEDULED_FOOT)
        table(("Setting","What it controls"),_CLOUD,(1.5,4.5))

# What is coming to Projects — classic edition only (None in threads).
_UP=getattr(C,"UPCOMING_THREADS",None)
if _UP:
    h2("Coming soon: the new Projects")
    p(_UP["intro"])
    table(("Feature","What it does"),_UP["features"],(1.7,4.3))
    h3("What changes in the labs")
    table(_UP["mapping_head"],_UP["mapping_rows"],(3.0,3.0),bold_first=False)
    p(_UP["foot"])

# ---------------- topic theory (placed before the lab that uses it) ----------------
def t1_family():
    h2("The Claude family")
    table(("Product","Role","What it does"),
          [(n,r,list(its)) for n,r,its in C.CLAUDE_THREE],   # a list renders as bullets
          (1.5,1.5,3.0))
    p(C.CLAUDE_THREE_FOOT)
def t1_research():
    if C.THREADS:
        h2("Market research in parallel threads")
        p("In Lab 2 you hand the lead agent one research job with four independent subjects. It "
          "starts one thread per subject:")
        table(("Thread","What it researches"),C.RESEARCH_THREADS,(1.6,4.4))
    else:
        h2("Market research in parallel tasks")
        p("In Lab 2 you start four research tasks in the project at once, one per independent "
          "subject, then a fifth task merges them:")
        table(("Task","What it researches"),C.RESEARCH_THREADS,(1.6,4.4))
    p(C.RESEARCH_FOOT)
def t1_artefacts():
    h2("Live Artefacts")
    p("A Live Artefact is an interactive tool Claude builds and keeps in your workspace. You make one in five moves:")
    defs(C.LIVE_ARTEFACT_FLOW,ordered=True)
    h3("What an artefact can and cannot do")
    compare(C.ARTEFACT_LIMITS)
def t2_connectors():
    h2("Connectors and the context window")
    p("A connector lets Claude reach a service outside the chat. These are the ones this course uses:")
    table(("Connector","What Claude can do"),C.CONNECTOR_CAPABILITIES,(1.6,4.4))
    h3("Keeping the context window lean")
    p(C.CONTEXT_WINDOW_QUOTE + " Five levers decide what goes into it:")
    table(("Lever","What it adds"),C.CONTEXT_LEVERS,(1.7,4.3))
    p(C.CONTEXT_LEVERS_FOOT)
    p("A window that fills up costs you in four ways:")
    table(("Problem","What happens"),C.CONTEXT_ROT,(1.7,4.3))
    p(C.CONTEXT_ROT_FOOT)
def t2_slides():
    h2("Slides, Docs and built-in skills")
    table(("Format","What it is"),C.SLIDES_DOCS,(1.5,4.5))
    p(C.SLIDES_DOCS_FOOT)
def t2_skills():
    h2("Agent Skills")
    p("An Agent Skill is a folder of instructions Claude loads on demand, so a routine runs the "
      "same way every time. A skill has three layers, and only the first is always in context:")
    table(("Layer","What it holds","When Claude loads it"),C.SKILL_LAYERS,(1.2,2.2,2.6))
    h3("Do, Verify, Save, Update")
    p("You never write a skill from a blank page. You do the task once, then save the verified routine:")
    defs(C.SKILL_PROCESS,ordered=True)
def t2_plugins():
    h2("Skills, plugins and mods")
    p(C.PLUGIN_FOOT)
    table(C.SKILLS_PLUGINS_MODS_HEAD,C.SKILLS_PLUGINS_MODS_ROWS,(1.2,1.6,1.6,1.6))
    h3("Mods")
    defs(C.MODS)
def t3_social():
    if C.THREADS:
        h2("Social content in parallel threads")
        p("Each platform has its own format, so Lab 8 gives every platform its own thread:")
    else:
        h2("Social content for four platforms")
        p("Each platform has its own format, so the Lab 8 prompt gives every platform its own spec:")
    table(("Platform","Format and rules"),C.SOCIAL_SPECS,(1.4,4.6))
    p(C.SOCIAL_SPECS_FOOT)
    h3("From brief to calendar")
    defs(C.CONTENT_THREADS_FLOW,ordered=True)
def t3_video():
    h2("The video pipeline")
    p("Claude does not edit video by hand — it writes and runs the code that renders it. The work runs in five stages:")
    defs(C.VIDEO_PIPELINE,ordered=True)
    h3("Video specifications")
    table(("Spec","Requirement"),C.VIDEO_SPECS,(1.3,4.7))
def t3_scheduled():
    h2("Scheduled tasks")
    p("A scheduled task has six parts. Write each one down before you schedule anything:")
    table(("Part","What to decide"),C.SCHEDULED_ANATOMY,(1.4,4.6))
    h3("Examples for the launch")
    table(("Schedule","Task"),C.SCHEDULED_EXAMPLES,(2.0,4.0))
    p(C.SCHEDULED_FOOT)
    h3("Responsible marketing with Claude")
    p("Speed is only useful if what goes out is true and fair. Keep these rules on every job:")
    table(("Rule","In practice"),C.RESPONSIBLE_MARKETING,(1.7,4.3))

# topic number -> {lab number: [sections placed immediately before that lab]}
THEORY={
    1:{1:[t1_family],2:[t1_research],3:[t1_artefacts]},
    2:{4:[t2_connectors],5:[t2_slides],6:[t2_skills],7:[t2_plugins]},
    3:{8:[t3_social],9:[t3_video],10:[t3_scheduled]},
}

# ---------------- one lab, in full ----------------
def lab(a):
    folder=lab_folder(a)
    h2(f"Lab {a['num']} — {a['title']}")
    figures_for(f"lab{a['num']}")
    lead("Objective",f"By the end of this lab you can {a['objective'].rstrip('.')}.")
    lead("Scenario",a["desc"])
    lead("You will build",a["build"])
    lead("Tools",a["services"]+".")
    h3("Lab files")
    files=[f"Prompt PDF: {C.LABS_SUBDIR}/{folder}/{folder}.pdf"]
    files+=[f"{C.LABS_SUBDIR}/{folder}/data/{d}" for d in a.get("data",[])]
    bullets(files)
    if not a.get("data"):
        p("This lab has no data folder — it works from what the project already holds.")
    if a.get("tools_sections"):
        h3("Tips")
        for title,body in a["tools_sections"]:
            h4(title); p(_oneline(body))
    # Each prompt is shown directly under the step where it is sent
    # (prompt_step / block["step"], 1-based into steps).
    n=len(a["steps"]); at={i:[] for i in range(1,n+1)}; loose=[]
    if a.get("prompt_blocks"):
        for b_ in a["prompt_blocks"]:
            item=dict(label=b_["label"],lead=b_.get("lead_lg",b_.get("lead")),
                      text=_wrap_prompt(b_["text"]),after=b_.get("after_lg",b_.get("after")))
            (at[b_["step"]] if b_.get("step") in at else loose).append(item)
    elif a.get("prompt"):
        item=dict(label=None,lead=None,text=_wrap_prompt(a["prompt"]),after=None,
                  note=a.get("prompt_note"))
        (at[a["prompt_step"]] if a.get("prompt_step") in at else loose).append(item)
    h3("Step-by-step")
    if a.get("prompt_blocks"):
        p(a.get("prompt_intro_lg",a.get("prompt_intro","")).replace("**",""))
        if a.get("prompt_note"): note(a["prompt_note"])
    for n_ in a.get("prompt_notes",[]): note(n_)
    for it in loose:                  # a prompt with no step recorded: before the steps
        if it.get("label"): h4(it["label"])
        if it.get("lead"): p(it["lead"])
        prompt(it["text"])
        if it.get("after"): p(it["after"])
        if it.get("note"): note(it["note"])
    steps([(si,sc,at[i]) for i,(si,sc) in enumerate(a["steps"],1)])
    h3("Check")
    check(a["test"])

for t in C.TOPICS:
    h1(f"Topic {t['code']} — {t['title']}",newpage=True)
    p(t["subtitle"])
    figures_for(f"topic{t['num']}")
    h3("Key concepts")
    table(("Concept","What it means"),t["concepts"],(1.7,4.3))
    for i_,a in enumerate([x for x in ACT if x["topic"]==t["num"]]):
        # Each lab unit (its theory, then the lab) starts on a fresh page, so a
        # lab can be printed or followed on its own. The first unit follows the
        # topic's key concepts directly.
        if i_: pagebreak()
        for fn in THEORY.get(t["num"],{}).get(a["num"],[]): fn()
        lab(a)

# ---------------- back matter ----------------
_WRAP=getattr(C,"LG_WRAPUP",None)
if _WRAP:
    h1(_WRAP.get("title","Wrap-Up"),newpage=True)
    if _WRAP.get("intro"): p(_WRAP["intro"])
    for sec in _WRAP.get("sections",[]):
        h3(sec.get("title",""))
        if sec.get("text"): p(sec["text"])
        if sec.get("bullets"): bullets(sec["bullets"])
    if getattr(C,"NEXT_STEPS",None):
        h3("Next steps"); defs(C.NEXT_STEPS)

_GLOSSARY=getattr(C,"LG_GLOSSARY",[])
if _GLOSSARY:
    h1("Glossary",newpage=True)
    table(("Term","Meaning"),_GLOSSARY,(1.6,4.4))

# ---------------- render Markdown ----------------
def _anchor(txt):
    return "".join(ch.lower() if ch.isalnum() else ("-" if ch in " -" else "") for ch in txt)

def _md(t):
    """Prose for Markdown: keep **bold**; put <PLACEHOLDERS> in code spans so a
    Markdown renderer does not swallow them as HTML tags."""
    return re.sub(r"(?<!`)<([A-Z][A-Z ]*)>(?!`)", r"`<\1>`", t)

def _cell(t):
    if isinstance(t,(list,tuple)): return "<br>".join("• "+_cell(x) for x in t)
    return _md(str(t)).replace("|","\\|").replace("\n"," ")

def render_md():
    out=[f"# {C.TITLE} — Learner Guide",""]
    out.append(f"**{C.EDITION_LABEL}**  |  **Course Code:** {C.COURSE_CODE}  |  **Conducted by:** {C.ORG} ({C.UEN.replace('UEN: ','UEN ')})  |  **Version {C.VERSION} · {C.VERSION_DATE}**")
    out.append("")
    out.append("## Contents"); out.append("")
    for kind,*rest in B:
        if kind=="h1": out.append(f"- [{rest[0]}](#{_anchor(rest[0])})")
        elif kind=="h2": out.append(f"  - [{rest[0]}](#{_anchor(rest[0])})")
    out.append("")
    for kind,*rest in B:
        if kind=="h1": out+=["",f"## {rest[0]}",""]
        elif kind=="h2": out+=["",f"### {rest[0]}",""]
        elif kind=="h3": out+=[f"#### {rest[0]}",""]
        elif kind=="h4": out+=[f"**{rest[0]}**",""]
        elif kind=="p": out+=[_md(rest[0]),""]
        elif kind=="lead": out+=[f"**{rest[0]}:** {_md(rest[1])}",""]
        elif kind=="bullets": out+=[f"- {_md(x)}" for x in rest[0]]+[""]
        elif kind=="defs":
            pairs,ordered=rest
            for i,(term,defn) in enumerate(pairs,1):
                out.append(f"{f'{i}.' if ordered else '-'} **{term}** — {_md(defn)}")
            out.append("")
        elif kind=="steps":
            for i,(instr,why,att) in enumerate(rest[0],1):
                out.append(f"{i}. {_md(instr)}" + ("  " if why else ""))   # hard break before Why
                if why: out.append(f"   *Why:* {_md(why)}")
                for it in att:        # indented 3 spaces so it stays inside the list item
                    out.append("")
                    if it.get("label"): out+=[f"   **{it['label']}**",""]
                    if it.get("lead"): out+=[f"   {_md(it['lead'])}",""]
                    out+=["   *PROMPT — copy into Claude*","","   ```text"]
                    out+=[("   "+l) if l else "" for l in it["text"].split("\n")]
                    out+=["   ```",""]
                    if it.get("after"): out+=[f"   {_md(it['after'])}",""]
                    if it.get("note"): out+=[f"   > **Note:** {_md(it['note'])}",""]
                if att: out.append("")
            out.append("")
        elif kind=="prompt":
            out+=[f"*{rest[1]}*","","```text",rest[0],"```",""]
        elif kind=="note": out+=[f"> **Note:** {_md(rest[0])}",""]
        elif kind=="check": out+=[f"> **Check — you are done when:** {_md(rest[0])}",""]
        elif kind=="table":
            head,rows,_w,_b=rest
            out.append("| "+" | ".join(_cell(h) for h in head)+" |")
            out.append("|"+"|".join("---" for _ in head)+"|")
            for r in rows: out.append("| "+" | ".join(_cell(c) for c in r)+" |")
            out.append("")
        elif kind=="figure":
            _path,fname,num,title,caption=rest
            out+=[f"![Figure {num} — {title}]({SHOT_REL}/{fname})","",
                  f"*Figure {num} — {title}. {caption}*",""]
        elif kind in ("rule","pagebreak"): out+=["---",""] if kind=="rule" else []
    return "\n".join(out).rstrip()+"\n"

MD_OUT=os.path.join(REPO,f"LG-{C.SHORT_TITLE}.md")
with open(MD_OUT,"w") as f: f.write(render_md())
print("Saved",MD_OUT)

# ---------------- render DOCX ----------------
BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
CLAY=RGBColor(0xD9,0x77,0x57); WHITE=RGBColor(0xFF,0xFF,0xFF)
PROMPT_FILL="FBF3EF"; CHECK_FILL="E8F7EE"; HEAD_FILL="1F6FEB"; ZEBRA_FILL="F3F5F8"
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)
prodoc.add_cover_page(doc,"LEARNER GUIDE",C.TITLE,f"{C.VERSION.lstrip('v')} · {C.VERSION_DATE}",
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE, edition=C.EDITION_LABEL)
prodoc.add_version_control(doc,getattr(C,"VERSION_HISTORY",
 [(C.VERSION.lstrip("v"),C.VERSION_DATE,"Initial release.",C.TRAINER)]))
prodoc.add_toc(doc)

def _runs(para,text,size=None,color=None,bold=False):
    """Add `text` to `para`, turning Markdown **bold** into real bold runs."""
    for part in re.split(r"(\*\*.+?\*\*)",text):
        if not part: continue
        b=part.startswith("**") and part.endswith("**") and len(part)>4
        r=para.add_run(part[2:-2] if b else part); r.bold=b or bold
        if size: r.font.size=Pt(size)
        if color is not None: r.font.color.rgb=color
    return para

def _keep_next(para,on=True):
    para.paragraph_format.keep_with_next=on; return para

def _no_split(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement("w:cantSplit"); e.set(qn("w:val"),"true"); trPr.append(e)

def _repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); trPr.append(e)

def _widths(tbl,widths): prodoc.set_col_widths(tbl,widths)

def _box(fill,width=6.0):
    """One-cell shaded table that never splits across a page."""
    tbl=doc.add_table(rows=1,cols=1); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=tbl.cell(0,0); prodoc._shade_cell(cell,fill); _widths(tbl,(width,)); _no_split(tbl.rows[0])
    return cell

def _prompt_box(text,header="PROMPT — copy into Claude",indent=0):
    """Shaded one-cell table: keeps the prompt visually separate from the prose,
    stops Word reflowing the line breaks the learner must keep, and (cantSplit)
    never breaks a prompt across two pages. `indent` (pt) aligns it under a step."""
    if len(doc.paragraphs): _keep_next(doc.paragraphs[-1])
    cell=_box(PROMPT_FILL,width=6.0-indent/72.0)
    if indent:
        tblPr=cell._tc.getparent().getparent().tblPr
        jc=tblPr.find(qn("w:jc"))
        if jc is not None: tblPr.remove(jc)
        ind=OxmlElement("w:tblInd"); ind.set(qn("w:w"),str(int(indent*20))); ind.set(qn("w:type"),"dxa")
        tblPr.append(ind)
    hdr=cell.paragraphs[0]; hdr.paragraph_format.space_after=Pt(3); _keep_next(hdr)
    hr=hdr.add_run(header); hr.bold=True; hr.font.size=Pt(8.5); hr.font.color.rgb=CLAY
    lines=text.split("\n")
    for li,line in enumerate(lines):
        para=cell.add_paragraph(); _keep_next(para)
        para.paragraph_format.space_after=Pt(4 if li==len(lines)-1 else 0)
        r=para.add_run(line if line else " ")
        r.font.name="Courier New"; r.font.size=Pt(9.5); r.font.color.rgb=DARK
    _gap(6)

def _ipara(indent,space_after=4):
    para=doc.add_paragraph(); para.paragraph_format.left_indent=Pt(indent)
    para.paragraph_format.space_after=Pt(space_after); return para

def _gap(pt=4):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(0)
    para.paragraph_format.space_before=Pt(0)
    r=para.add_run(""); r.font.size=Pt(pt); para.paragraph_format.line_spacing=Pt(pt)

def _figure(path,num,title,caption):
    try:
        from PIL import Image
        w,h=Image.open(path).size; width=min(6.0, 4.5*w/h)   # at most 4.5in tall
    except Exception:
        width=6.0
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; _keep_next(para)
    para.add_run().add_picture(path,width=Inches(width))
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after=Pt(10)
    r=cap.add_run(f"Figure {num} — {title}. "); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=DARK
    r=cap.add_run(caption); r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY

_break_next=False
for kind,*rest in B:
    if kind=="h1":
        hp=doc.add_heading(rest[0],level=1)
        if rest[1]: hp.paragraph_format.page_break_before=True
    elif kind=="pagebreak": _break_next=True       # applied to the next heading
    elif kind=="h2":
        hp=doc.add_heading(rest[0],level=2)
        if _break_next: hp.paragraph_format.page_break_before=True; _break_next=False
    elif kind=="h3":
        para=_keep_next(doc.add_paragraph()); para.paragraph_format.space_before=Pt(8)
        para.paragraph_format.space_after=Pt(4)
        r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BRAND
    elif kind=="h4":
        para=_keep_next(doc.add_paragraph()); para.paragraph_format.space_before=Pt(6)
        para.paragraph_format.space_after=Pt(2)
        r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=DARK
    elif kind=="p": _runs(doc.add_paragraph(),rest[0])
    elif kind=="lead":
        para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(4)
        r=para.add_run(rest[0]+": "); r.bold=True; r.font.color.rgb=DARK
        _runs(para,rest[1])
    elif kind=="bullets":
        for x in rest[0]: _runs(doc.add_paragraph(style="List Bullet"),x)
    elif kind=="defs":
        pairs,ordered=rest
        for i,(term,defn) in enumerate(pairs,1):
            if ordered:
                para=doc.add_paragraph(); para.paragraph_format.left_indent=Pt(18)
                para.paragraph_format.first_line_indent=Pt(-18); para.paragraph_format.space_after=Pt(3)
                r=para.add_run(f"{i}.\t"); r.bold=True; r.font.color.rgb=BRAND
            else:
                para=doc.add_paragraph(style="List Bullet")
            r=para.add_run(term+" — "); r.bold=True; _runs(para,defn)
    elif kind=="steps":
        for i,(instr,why,att) in enumerate(rest[0],1):
            # Literal per-lab step numbers (Word's List Number style numbers
            # continuously across the whole document — 1..N must restart per lab).
            para=doc.add_paragraph()
            para.paragraph_format.left_indent=Pt(22); para.paragraph_format.first_line_indent=Pt(-22)
            para.paragraph_format.space_after=Pt(2 if why else 4)
            r=para.add_run(f"{i}.\t"); r.bold=True; r.font.color.rgb=BRAND
            _runs(para,instr)
            if why:
                _keep_next(para)
                wp=doc.add_paragraph(); wp.paragraph_format.left_indent=Pt(22)
                wp.paragraph_format.space_after=Pt(6)
                r=wp.add_run("Why: "); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
                _runs(wp,why,size=9.5,color=GREY)
            for it in att:
                if it.get("label"):
                    lp=_keep_next(_ipara(22,2)); lp.paragraph_format.space_before=Pt(4)
                    r=lp.add_run(it["label"]); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=DARK
                if it.get("lead"): _runs(_keep_next(_ipara(22,3)),it["lead"])
                _prompt_box(it["text"],indent=22)
                if it.get("after"): _runs(_ipara(22,8),it["after"])
                if it.get("note"):
                    np_=_ipara(22,8); r=np_.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
                    _runs(np_,it["note"],size=10)
    elif kind=="prompt":
        _prompt_box(*rest)
    elif kind=="check":
        if len(doc.paragraphs): _keep_next(doc.paragraphs[-1])
        cell=_box(CHECK_FILL)
        para=cell.paragraphs[0]; para.paragraph_format.space_after=Pt(2)
        r=para.add_run("You are done when: "); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x1E,0x7B,0x46)
        _runs(para,rest[0],size=10)
        _gap(6)
    elif kind=="note":
        para=doc.add_paragraph(); r=para.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
        _runs(para,rest[0],size=10)
    elif kind=="table":
        head,rows,widths,bold_first=rest
        if len(doc.paragraphs): _keep_next(doc.paragraphs[-1])
        small=len(rows)<=10          # short tables stay on one page
        tbl=doc.add_table(rows=0,cols=len(head)); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
        hr_=tbl.add_row(); _repeat_header(hr_); _no_split(hr_)
        for c,h in zip(hr_.cells,head):
            c.text=""; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=WHITE
            c.paragraphs[0].paragraph_format.space_before=Pt(1.5); c.paragraphs[0].paragraph_format.space_after=Pt(1.5)
            _keep_next(c.paragraphs[0])
            prodoc._shade_cell(c,HEAD_FILL)
        for ri,row in enumerate(rows):
            tr=tbl.add_row(); _no_split(tr)
            for ci,(c,val) in enumerate(zip(tr.cells,row)):
                c.text=""; cp=c.paragraphs[0]
                items=val if isinstance(val,(list,tuple)) else [val]
                for k_,item in enumerate(items):
                    if k_: cp=c.add_paragraph()
                    cp.paragraph_format.space_before=Pt(1.5); cp.paragraph_format.space_after=Pt(1.5)
                    if isinstance(val,(list,tuple)):     # a bulleted list inside the cell
                        cp.paragraph_format.left_indent=Pt(10); cp.paragraph_format.first_line_indent=Pt(-10)
                        cp.add_run("•\t").font.size=Pt(9.5)
                        cp.paragraph_format.tab_stops.add_tab_stop(Pt(10))
                    _runs(cp,str(item),size=9.5,bold=(ci==0 and bold_first))
                    if small and ri<len(rows)-1: _keep_next(cp)
                if ri%2==1: prodoc._shade_cell(c,ZEBRA_FILL)
        _widths(tbl,widths)
        _gap(6)
    elif kind=="figure":
        _path,_fname,num,title,caption=rest
        _figure(_path,num,title,caption)
    elif kind=="rule": doc.add_paragraph("")

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
DOCX_OUT=os.path.join(REPO,C.COURSEWARE_SUBDIR,f"LG-{C.SHORT_TITLE}.docx")   # courseware/ or courseware/threads-edition/
os.makedirs(os.path.dirname(DOCX_OUT),exist_ok=True)
prodoc.set_properties(doc, f"Learner Guide — {C.TITLE}", C.COURSE_CODE)
doc.save(DOCX_OUT)
print("Saved",DOCX_OUT)
