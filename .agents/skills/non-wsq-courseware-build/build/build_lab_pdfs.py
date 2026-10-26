#!/usr/bin/env python3
"""Build each lab's prompt PDF straight from the lab data.

The PDF is the one document in a lab folder, read top to bottom: course and
lab header, the course scenario in a short box, the lab's Scenario (desc) and
Objective, what you build, the data files, any "before you start" notes, then
ONE numbered list of steps with each prompt box — label, lead line, the prompt
in a shaded Courier box, and its after line — sitting directly under the step
it belongs to (step → prompt → step → prompt), and finally the check.

Rendered with python-docx + LibreOffice (already required by the rest of the
pipeline), so no extra dependency is introduced. The intermediate DOCX files
are written to REPO/.lo-labpdf/ (gitignored by the .lo-*/ rule) and are simply
overwritten on the next run — the lab folders stay PDF-only.
"""
import os, re, sys, subprocess, pathlib

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
import prodoc
from build_labs import ACTS, REPO, TOPIC, lab_dir, lab_folder, pdf_name, blocks, steps_with_prompts

# Claude Style palette (claude_style.py), as docx colours.
CLAY = RGBColor(0xD9, 0x77, 0x57)
INK = RGBColor(0x21, 0x20, 0x1C)
MUTED = RGBColor(0x6B, 0x68, 0x62)
GREEN = RGBColor(0x4A, 0x8F, 0x5F)
BOX_FILL = "FBF3EF"       # warm tint behind each prompt
NOTE_FILL = "F5F3EE"
HEAD_FONT = "Georgia"     # Cambria is not installed for LibreOffice on every machine
BODY_FONT = "Arial"
CODE_FONT = "Courier New"
TEXT_W = 6.67             # A4 width less two 0.8in margins
STEP_IN = 0.3             # hanging indent of the numbered steps

SOFFICE = os.environ.get("SOFFICE", "soffice")
WORK = os.path.join(REPO, ".lo-labpdf")


# ------------------------------------------------------------------ primitives
def _para(doc_or_cell, text="", size=10.5, color=INK, bold=False, italic=False,
          before=0, after=6, font=None, keep_next=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.keep_with_next = keep_next
    if text:
        _inline(p, text, size=size, color=color, bold=bold, italic=italic, font=font)
    return p


def _inline(p, text, size=10.5, color=INK, bold=False, italic=False, font=None):
    """Render **bold** and `code` spans from the Markdown-flavoured source."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        b, f = bold, font
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, f = part[1:-1], CODE_FONT
        r = p.add_run(part)
        r.bold, r.italic = b, italic
        r.font.size, r.font.color.rgb = Pt(size), color
        if f:
            r.font.name = f
    return p


def _rule(p, color="D97757", size=12):
    """A coloured bottom border under a paragraph — the clay header rule."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "4"), ("color", color)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _no_split(table):
    """Ask Word/LibreOffice not to break a table row across pages. A prompt box
    is one row, so this keeps every prompt on a single page where it fits."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)


def _box(doc, fill, border="E4E1D9", indent=0.0):
    """A one-cell table used as a shaded panel; returns the cell. `indent`
    (inches) shifts it right and narrows it, so a prompt box lines up under
    the text of the step it belongs to instead of under the step number."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(TEXT_W - indent)
    tbl.columns[0].width = Inches(TEXT_W - indent)
    if indent:
        tblPr = tbl._tbl.tblPr
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(indent * 1440)))
        ind.set(qn("w:type"), "dxa")
        tblPr.append(ind)
    prodoc._shade_cell(cell, fill)
    # Soften the Table Grid border to the hairline colour.
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "18" if side == "left" else "4")
        el.set(qn("w:color"), "D97757" if side == "left" else border)
        borders.append(el)
    tcPr.append(borders)
    _no_split(tbl)
    return cell


def heading(doc, text, before=12):
    p = _para(doc, before=before, after=4, keep_next=True)
    r = p.add_run(text)
    r.bold, r.font.name = True, HEAD_FONT
    r.font.size, r.font.color.rgb = Pt(13.5), INK
    return p


def label(doc, text, color=CLAY, before=8):
    p = _para(doc, before=before, after=2, keep_next=True)
    r = p.add_run(text.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), color
    return p


# ------------------------------------------------------------------ blocks
def prompt_box(doc, text, num, title, indent=0.0):
    """The prompt itself: a shaded Courier box, one paragraph per source line
    so the line breaks survive copy and paste. <PLACEHOLDERS> are picked out."""
    cell = _box(doc, BOX_FILL, indent=indent)
    hdr = cell.paragraphs[0]
    hdr.paragraph_format.space_after = Pt(4)
    hdr.paragraph_format.keep_with_next = True
    r = hdr.add_run(f"LAB {num:02d} · {title.upper()}   ·   COPY INTO CLAUDE")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8), CLAY
    for line in text.strip("\n").split("\n"):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        for part in re.split(r"(<[A-Z][^>]*>)", line) or [""]:
            run = p.add_run(part)
            run.font.name, run.font.size = CODE_FONT, Pt(9)
            if part.startswith("<") and part.endswith(">"):
                run.bold, run.font.color.rgb = True, CLAY
            else:
                run.font.color.rgb = INK
    _para(doc, after=2, size=4)


def note(doc, text, title="Note"):
    cell = _box(doc, NOTE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{title}: ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), CLAY
    _inline(p, " ".join(text.split()), size=9.5, color=INK)
    _para(doc, after=2, size=4)


def scenario_box(doc):
    """The overarching course story, short, on page 1 of every lab PDF — so a
    learner who opens Lab 7 cold still knows whose launch this is."""
    sc = C.COURSE_SCENARIO
    cell = _box(doc, NOTE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"THE COURSE SCENARIO  ·  {sc['title'].upper()}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8), CLAY
    for head, key in (("The situation", "situation"), ("Your role", "role"), ("The mission", "mission")):
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        _inline(q, f"**{head}.** {sc[key]}", size=9, color=INK)
    _para(doc, after=0, size=4)


# ------------------------------------------------------------------ document
def lab_docx(a, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)        # A4
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(10.5)

    t = TOPIC[a["topic"]]
    num = a["num"]

    # Header: course eyebrow, lab title, topic line, clay rule.
    _para(doc, f"{C.TITLE.upper()}  ·  {C.COURSE_CODE}  ·  {C.VERSION}  ·  {C.VERSION_DATE}",
          size=8.5, color=CLAY, bold=True, after=2)
    p = _para(doc, after=2)
    r = p.add_run(f"Lab {num:02d} — {a['title']}")
    r.bold, r.font.name, r.font.size, r.font.color.rgb = True, HEAD_FONT, Pt(20), INK
    p = _para(doc, f"Topic {t['code']} · {t['title']}", size=10, color=MUTED, after=8)
    _rule(p)

    scenario_box(doc)

    heading(doc, "Scenario", before=8)
    _para(doc, a["desc"])
    heading(doc, "Objective")
    _para(doc, f"By the end of this lab you can {a['objective']}.")
    heading(doc, "You will build")
    _para(doc, a["build"], after=3)
    _para(doc, f"**Tools:** {a['services']}", size=9.5, color=MUTED)

    heading(doc, "Data files")
    if a.get("data"):
        _para(doc, f"From this lab's **data/** folder in {lab_folder(a)}:", after=3, keep_next=True)
        for f in a["data"]:
            p = _para(doc, after=1)
            p.paragraph_format.left_indent = Inches(0.2)
            _inline(p, "•  ", color=CLAY)
            _inline(p, f"data/{f}", size=9.5, font=CODE_FONT)
    else:
        _para(doc, "None — this lab works from what the project already holds.")

    # ---- before you start: the prompt note and any tool notes are
    # prerequisites (attach these files, install the plugin, swap in your email)
    # so they come before the walk-through, not after it.
    notes = list(a.get("prompt_notes", [])) + ([a["prompt_note"]] if a.get("prompt_note") else [])
    tools = a.get("tools_sections") or []
    if notes or tools:
        heading(doc, "Before you start")
        for n in notes:
            note(doc, n)
        for head, body in tools:
            note(doc, body, title=head)

    # ---- the walk-through: step → prompt → step → prompt
    heading(doc, f"Steps and prompts — {a['prompt_title']}", before=14)
    _para(doc, f"Each prompt sits under the step that sends it. Copy it from the box, or with one "
               f"click from the Lab Prompts page: {C.PROMPTS_URL}", size=9.5, color=MUTED)
    if a.get("prompt_intro"):
        _para(doc, a["prompt_intro"])
    multi = len(blocks(a)) > 1
    for i, step, why, bl in steps_with_prompts(a):
        p = _para(doc, after=1, before=5, keep_next=True)
        p.paragraph_format.left_indent = Inches(STEP_IN)
        p.paragraph_format.first_line_indent = Inches(-STEP_IN)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(STEP_IN))
        _inline(p, f"{i}.\t", bold=True, color=CLAY)
        _inline(p, step, bold=bool(bl))
        if why:
            w = _para(doc, why, size=9, color=MUTED, after=3, keep_next=bool(bl))
            w.paragraph_format.left_indent = Inches(STEP_IN)
        for b in bl:
            if multi:
                lp = label(doc, b["label"], before=4)
                lp.paragraph_format.left_indent = Inches(STEP_IN)
            if b.get("lead"):
                lp = _para(doc, b["lead"], size=10, after=3, keep_next=True)
                lp.paragraph_format.left_indent = Inches(STEP_IN)
            prompt_box(doc, b["text"], num, b["label"] if multi else "Prompt", indent=STEP_IN)
            if b.get("after"):
                ap = _para(doc, b["after"], size=10, after=6)
                ap.paragraph_format.left_indent = Inches(STEP_IN)
            for n in b.get("notes", []):
                note(doc, n)

    heading(doc, "Check", before=10)
    cell = _box(doc, "EEF4EF")
    p = cell.paragraphs[0]
    r = p.add_run("DONE WHEN  ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), GREEN
    _inline(p, a["test"], size=10)

    _para(doc, C.LAB_NOTE, size=8.5, color=MUTED, italic=True, before=10)

    prodoc.add_page_numbers(doc)
    prodoc.set_properties(doc, f"Lab {num:02d} — {a['title']}", C.TITLE)
    doc.save(path)


def render(docx_path, outdir):
    """DOCX → PDF with a private LibreOffice profile, so a soffice instance
    another build is already running cannot swallow the conversion."""
    # as_uri() percent-encodes the spaces in the repo path; a raw "file://"
    # path with spaces makes soffice abort.
    profile = pathlib.Path(WORK, "lo-profile").as_uri()
    subprocess.run([SOFFICE, f"-env:UserInstallation={profile}", "--headless",
                    "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def main():
    os.makedirs(WORK, exist_ok=True)
    made = 0
    for a in ACTS:
        d = lab_dir(a)
        os.makedirs(d, exist_ok=True)
        docx = os.path.join(WORK, f"{lab_folder(a)}.docx")
        lab_docx(a, docx)
        before = os.path.getmtime(os.path.join(d, pdf_name(a))) if os.path.exists(os.path.join(d, pdf_name(a))) else 0
        pdf = render(docx, d)
        ok = os.path.exists(pdf) and os.path.getmtime(pdf) > before
        made += ok
        print(f"  {'ok  ' if ok else 'FAIL'} labs/{lab_folder(a)}/{pdf_name(a)}")
    print(f"  {made}/{len(ACTS)} lab PDFs")
    if made != len(ACTS):
        sys.exit(1)


if __name__ == "__main__":
    print("Rendering lab prompt PDFs")
    main()
